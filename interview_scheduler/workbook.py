"""Read the input workbook and write the finished interview schedule."""

from __future__ import annotations

import os
import re
import tempfile
import zipfile
from datetime import datetime, time
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from interview_scheduler.errors import InputValidationError
from interview_scheduler.models import (
    AvailabilityGrid,
    Schedule,
    TimeSlot,
    WeekdayAvailability,
)

WEEKDAYS = ("Monday", "Tuesday", "Wednesday", "Thursday", "Friday")
EXPECTED_TIME_SLOTS = tuple(TimeSlot(start, start + 30) for start in range(600, 1020, 30))
_TIME_RANGE = re.compile(
    r"^(?P<start>\d{1,2}:\d{2})\s*(?P<start_meridiem>AM|PM|NN)?\s*-\s*"
    r"(?P<end>\d{1,2}:\d{2})\s*(?P<end_meridiem>AM|PM|NN)$",
    re.IGNORECASE,
)
_MAX_INPUT_BYTES = 50 * 1024 * 1024
_MAX_ARCHIVE_ENTRIES = 2_000
_MAX_UNCOMPRESSED_BYTES = 200 * 1024 * 1024
_MAX_ENTRY_BYTES = 100 * 1024 * 1024
_MAX_COMPRESSION_RATIO = 200
_REQUIRED_XLSX_PARTS = frozenset({"[Content_Types].xml", "_rels/.rels", "xl/workbook.xml"})


def _clock_minutes(clock_text: str, meridiem: str) -> int:
    hour_text, minute_text = clock_text.split(":", maxsplit=1)
    hour, minute = int(hour_text), int(minute_text)
    if hour not in range(1, 13) or minute not in range(60):
        raise InputValidationError(f"Invalid clock time: {clock_text} {meridiem}.")
    marker = meridiem.upper()
    if marker == "NN":
        if (hour, minute) != (12, 0):
            raise InputValidationError("The 'NN' marker is only valid for exactly 12:00 noon.")
        return 720
    return (hour % 12) * 60 + minute + (720 if marker == "PM" else 0)


def parse_time_slot(value: object) -> TimeSlot:
    """Parse a textual Excel time range into a validated half-hour slot."""

    if not isinstance(value, str):
        raise InputValidationError("Time cells must contain ranges such as 10:00-10:30 AM.")
    normalized = re.sub(r"[–—]", "-", value.strip())
    match = _TIME_RANGE.fullmatch(normalized)
    if match is None:
        raise InputValidationError(f"Invalid time range: {value!r}.")

    end_meridiem = match.group("end_meridiem").upper()
    end_minutes = _clock_minutes(match.group("end"), end_meridiem)
    start_marker = match.group("start_meridiem")
    if start_marker:
        start_minutes = _clock_minutes(match.group("start"), start_marker.upper())
    else:
        candidates = (
            _clock_minutes(match.group("start"), "AM"),
            _clock_minutes(match.group("start"), "PM"),
        )
        matching = tuple(candidate for candidate in candidates if end_minutes - candidate == 30)
        if not matching:
            raise InputValidationError(f"Time range {value!r} must be a 30-minute interval.")
        if len(matching) != 1:
            raise InputValidationError(f"Could not infer the starting AM/PM marker in {value!r}.")
        start_minutes = matching[0]
    try:
        return TimeSlot(start_minutes, end_minutes)
    except ValueError as error:
        raise InputValidationError(f"Time range {value!r} must be a 30-minute interval.") from error


def _normalized_name(value: object, location: str) -> str:
    if not isinstance(value, str):
        raise InputValidationError(f"{location} must contain an interviewer name.")
    name = " ".join(value.split())
    if not name:
        raise InputValidationError(f"{location} must contain an interviewer name.")
    return name


def _status_is_available(value: object, location: str) -> bool:
    if not isinstance(value, str):
        raise InputValidationError(f"{location} must contain either 'available' or 'x'.")
    status = value.strip().casefold()
    if status == "available":
        return True
    if status == "x":
        return False
    raise InputValidationError(f"{location} contains {value!r}; expected 'available' or 'x'.")


def _reject_formula(cell: Cell, sheet_name: str) -> None:
    if cell.data_type == "f":
        raise InputValidationError(f"{sheet_name}!{cell.coordinate} contains a formula; enter a literal value instead.")


def _read_headers(worksheet) -> tuple[tuple[int, str], ...]:
    sheet_name = worksheet.title
    time_header = worksheet.cell(4, 1)
    _reject_formula(time_header, sheet_name)
    if not isinstance(time_header.value, str) or time_header.value.strip().casefold() != "time":
        raise InputValidationError(f"{sheet_name}!A4 must contain the header 'time'.")

    populated = tuple(column for column in range(2, 24) if worksheet.cell(4, column).value not in (None, ""))
    if len(populated) < 3:
        raise InputValidationError(f"{sheet_name} must contain at least three interviewer headers in B4:W4.")
    last_column = max(populated)
    if populated != tuple(range(2, last_column + 1)):
        raise InputValidationError(f"{sheet_name} has a blank interviewer header between B4 and {get_column_letter(last_column)}4.")

    headers = []
    for column in populated:
        cell = worksheet.cell(4, column)
        _reject_formula(cell, sheet_name)
        headers.append((column, _normalized_name(cell.value, f"{sheet_name}!{cell.coordinate}")))
    keys = tuple(name.casefold() for _, name in headers)
    if len(keys) != len(set(keys)):
        raise InputValidationError(f"{sheet_name} contains duplicate interviewer headers after normalization.")
    return tuple(headers)


def _validate_unused_columns(worksheet, active_columns: set[int]) -> None:
    for column in range(2, 24):
        if column in active_columns:
            continue
        for row in range(5, 19):
            cell = worksheet.cell(row, column)
            if cell.value not in (None, ""):
                raise InputValidationError(
                    f"{worksheet.title}!{cell.coordinate} has data below a blank interviewer header."
                )


def _read_weekday(worksheet) -> tuple[tuple[str, ...], tuple[tuple[TimeSlot, frozenset[str]], ...]]:
    headers = _read_headers(worksheet)
    _validate_unused_columns(worksheet, {column for column, _ in headers})
    rows = []
    for index, expected_slot in enumerate(EXPECTED_TIME_SLOTS, start=5):
        time_cell = worksheet.cell(index, 1)
        _reject_formula(time_cell, worksheet.title)
        parsed_slot = parse_time_slot(time_cell.value)
        if parsed_slot != expected_slot:
            raise InputValidationError(
                f"{worksheet.title}!A{index} must be {expected_slot.label}, not {parsed_slot.label}."
            )
        available = []
        for column, name in headers:
            cell = worksheet.cell(index, column)
            _reject_formula(cell, worksheet.title)
            if _status_is_available(cell.value, f"{worksheet.title}!{cell.coordinate}"):
                available.append(name)
        rows.append((parsed_slot, frozenset(available)))
    return tuple(name for _, name in headers), tuple(rows)


def _weekday_sheets(workbook) -> dict[str, object]:
    normalized = {" ".join(sheet.title.split()).casefold(): sheet for sheet in workbook.worksheets}
    missing = tuple(weekday for weekday in WEEKDAYS if weekday.casefold() not in normalized)
    if missing:
        raise InputValidationError(f"Missing required weekday sheet(s): {', '.join(missing)}.")
    return {weekday: normalized[weekday.casefold()] for weekday in WEEKDAYS}


def _canonicalize_rows(
    weekday: str,
    names: tuple[str, ...],
    rows: tuple[tuple[TimeSlot, frozenset[str]], ...],
    canonical_names: dict[str, str],
) -> WeekdayAvailability:
    name_keys = {name.casefold() for name in names}
    if name_keys != set(canonical_names):
        missing = sorted(set(canonical_names) - name_keys)
        extra = sorted(name_keys - set(canonical_names))
        detail = f"missing={missing or 'none'}, extra={extra or 'none'}"
        raise InputValidationError(f"{weekday} interviewer headers do not match Monday ({detail}).")
    canonical_rows = tuple(
        (slot, frozenset(canonical_names[name.casefold()] for name in available))
        for slot, available in rows
    )
    return WeekdayAvailability(weekday, canonical_rows)


def _preflight_xlsx_archive(workbook_path: Path) -> None:
    try:
        with zipfile.ZipFile(workbook_path) as archive:
            entries = tuple(entry for entry in archive.infolist() if not entry.is_dir())
    except zipfile.BadZipFile as error:
        raise InputValidationError("The input is not a valid .xlsx ZIP archive.") from error
    if len(entries) > _MAX_ARCHIVE_ENTRIES:
        raise InputValidationError(f"The input workbook contains too many ZIP entries ({len(entries)}).")
    names = {entry.filename for entry in entries}
    if not _REQUIRED_XLSX_PARTS.issubset(names):
        raise InputValidationError("The input archive is missing required .xlsx workbook parts.")
    if any(entry.flag_bits & 0x1 for entry in entries):
        raise InputValidationError("Encrypted .xlsx archives are not supported.")
    if any(entry.file_size > _MAX_ENTRY_BYTES for entry in entries):
        raise InputValidationError("An expanded workbook entry exceeds the 100 MB safety limit.")
    total_size = sum(entry.file_size for entry in entries)
    if total_size > _MAX_UNCOMPRESSED_BYTES:
        raise InputValidationError("The expanded workbook exceeds the 200 MB safety limit.")
    if any(entry.file_size / max(entry.compress_size, 1) > _MAX_COMPRESSION_RATIO for entry in entries):
        raise InputValidationError("The workbook contains a suspiciously compressed ZIP entry.")


def read_availability(path: str | Path) -> AvailabilityGrid:
    """Validate and normalize A4:W18 from the five weekday worksheets."""

    workbook_path = Path(path)
    if workbook_path.suffix.casefold() != ".xlsx":
        raise InputValidationError("The input must be an .xlsx workbook.")
    if not workbook_path.is_file():
        raise FileNotFoundError(f"Input workbook not found: {workbook_path}")
    if workbook_path.stat().st_size > _MAX_INPUT_BYTES:
        raise InputValidationError("The input workbook exceeds the 50 MB safety limit.")

    _preflight_xlsx_archive(workbook_path)
    workbook = load_workbook(workbook_path, read_only=True, data_only=False, keep_links=False)
    try:
        sheets = _weekday_sheets(workbook)
        raw = {weekday: _read_weekday(sheet) for weekday, sheet in sheets.items()}
    finally:
        workbook.close()

    monday_names = raw["Monday"][0]
    canonical_names = {name.casefold(): name for name in monday_names}
    sorted_names = tuple(sorted(canonical_names.values(), key=str.casefold))
    canonical_names = {name.casefold(): name for name in sorted_names}
    weekdays = tuple(
        _canonicalize_rows(weekday, *raw[weekday], canonical_names)
        for weekday in WEEKDAYS
    )
    return AvailabilityGrid(sorted_names, weekdays)


def _schedule_axes(schedule: Schedule) -> tuple[tuple, tuple[TimeSlot, ...]]:
    dates = tuple(sorted({assignment.interview_date for assignment in schedule.assignments}))
    slots = tuple(sorted({assignment.time_slot for assignment in schedule.assignments}))
    expected = {(interview_date, slot) for interview_date in dates for slot in slots}
    actual = {(assignment.interview_date, assignment.time_slot) for assignment in schedule.assignments}
    if actual != expected or len(actual) != len(schedule.assignments):
        raise ValueError("The schedule must contain exactly one assignment for every date/time combination.")
    return dates, slots


def _style_schedule(worksheet, date_count: int, slot_count: int) -> None:
    navy = PatternFill("solid", fgColor="17365D")
    pale_blue = PatternFill("solid", fgColor="DCE6F1")
    light_border = Border(bottom=Side(style="thin", color="CBD5E1"))
    header_range = worksheet.iter_rows(min_row=1, max_row=1, min_col=1, max_col=date_count + 1)
    for cell in next(header_range):
        cell.fill = navy
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in range(2, slot_count + 2):
        time_cell = worksheet.cell(row, 1)
        time_cell.fill = pale_blue
        time_cell.font = Font(bold=True, color="17365D")
        time_cell.alignment = Alignment(horizontal="center", vertical="center")
        for column in range(1, date_count + 2):
            worksheet.cell(row, column).border = light_border
        for column in range(2, date_count + 2):
            worksheet.cell(row, column).alignment = Alignment(vertical="center", wrap_text=True)
        worksheet.row_dimensions[row].height = 48
    worksheet.row_dimensions[1].height = 32
    worksheet.column_dimensions["A"].width = 20
    for column in range(2, date_count + 2):
        worksheet.column_dimensions[get_column_letter(column)].width = 24
    worksheet.freeze_panes = "B2"
    worksheet.sheet_view.showGridLines = False
    worksheet.page_setup.orientation = "landscape"
    worksheet.auto_filter.ref = f"A1:{get_column_letter(date_count + 1)}{slot_count + 1}"


def _populate_schedule(workbook: Workbook, schedule: Schedule) -> None:
    worksheet = workbook.active
    worksheet.title = "Interview Schedule"
    dates, slots = _schedule_axes(schedule)
    assignments = {
        (assignment.interview_date, assignment.time_slot): assignment.interviewers
        for assignment in schedule.assignments
    }
    worksheet.cell(1, 1, "Time")
    for column, interview_date in enumerate(dates, start=2):
        cell = worksheet.cell(1, column, datetime.combine(interview_date, time.min))
        cell.number_format = "mmmm d (dddd)"
    for row, slot in enumerate(slots, start=2):
        worksheet.cell(row, 1, slot.label)
        for column, interview_date in enumerate(dates, start=2):
            cell = worksheet.cell(row, column, "\n".join(assignments[(interview_date, slot)]))
            cell.data_type = "s"
    _style_schedule(worksheet, len(dates), len(slots))


def write_schedule(schedule: Schedule, output_path: str | Path) -> Path:
    """Write the schedule atomically so a failure cannot leave a partial workbook."""

    destination = Path(output_path)
    if destination.suffix.casefold() != ".xlsx":
        raise ValueError("The output path must end in .xlsx.")
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    _populate_schedule(workbook, schedule)
    descriptor, temporary_name = tempfile.mkstemp(
        dir=destination.parent,
        prefix=f".{destination.stem}-",
        suffix=".xlsx",
    )
    os.close(descriptor)
    temporary_path = Path(temporary_name)
    try:
        workbook.save(temporary_path)
        os.replace(temporary_path, destination)
    finally:
        workbook.close()
        temporary_path.unlink(missing_ok=True)
    return destination
